from pathlib import Path

from openpyxl import load_workbook

file_path = Path(__file__).resolve().parents[2] / "2027_BCI_Internship_Tracker.xlsx"
wb = load_workbook(str(file_path))
ws = wb["2.Research_Database"]

print("First row values:")
for col in range(1, 7):
    print(ws.cell(row=1, column=col).value)
