import hashlib
import hmac
import json
import logging

import openpyxl
import openpyxl.utils

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("ZeroHallucinationFailsafe")

VALID_STATUSES = (
    "VERIFIED_VALID",
    "VERIFIED_VALID_PAYWALLED",
    "VERIFIED_VALID_PORTAL_ZERO_OPENINGS",
    "VERIFIED_VALID_FACULTY_MIGRATED",
)


def verify_attestation_ledger_integrity(telemetry_payload_json: str, recorded_hash: str) -> bool:
    try:
        if not telemetry_payload_json or not recorded_hash:
            return False
        payload_dict = json.loads(telemetry_payload_json)
        canonical_jcs_bytes = json.dumps(
            payload_dict, sort_keys=True, separators=(",", ":"), ensure_ascii=False
        ).encode("utf-8")
        computed_hash = hashlib.sha256(canonical_jcs_bytes).hexdigest()
        return hmac.compare_digest(computed_hash.lower(), recorded_hash.lower())
    except Exception:
        return False


def enforce_zero_hallucination_excel_failsafe(excel_file_path: str):
    logger.info(f"Loading Excel file: {excel_file_path}")
    wb = openpyxl.load_workbook(excel_file_path)
    if "2.Research_Database" not in wb.sheetnames:
        raise ValueError("Worksheet '2.Research_Database' not found in target Excel file.")

    ws = wb["2.Research_Database"]
    headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]

    try:
        status_idx = headers.index("Verification_Status") + 1
        hash_idx = headers.index("Verification_Hash") + 1
        payload_idx = headers.index("Telemetry_Payload") + 1
    except ValueError as e:
        logger.error(f"Missing required metadata columns. Make sure columns were added. {e}")
        return

    stats = {
        "total_evaluated": 0,
        "visible_valid": 0,
        "hidden_unverified": 0,
        "quarantined_tampered": 0,
        "replay_attack_tampered": 0,
    }

    for row_idx in range(2, ws.max_row + 1):
        # Skip empty rows (Company is empty)
        if not ws.cell(row=row_idx, column=1).value:
            continue

        stats["total_evaluated"] += 1
        status_val = str(ws.cell(row=row_idx, column=status_idx).value or "").strip()
        hash_val = str(ws.cell(row=row_idx, column=hash_idx).value or "").strip()
        payload_val = str(ws.cell(row=row_idx, column=payload_idx).value or "").strip()

        if status_val in VALID_STATUSES:
            ledger_valid = len(hash_val) == 64 and verify_attestation_ledger_integrity(payload_val, hash_val)
            if not ledger_valid:
                ws.row_dimensions[row_idx].hidden = True
                ws.cell(row=row_idx, column=status_idx).value = "QUARANTINED_TAMPERED"
                stats["quarantined_tampered"] += 1
                stats["hidden_unverified"] += 1
                logger.warning(f"Row {row_idx}: Cryptographic attestation hash mismatch or missing. Row quarantined.")
                continue

            ws.row_dimensions[row_idx].hidden = False
            stats["visible_valid"] += 1
        else:
            # Legacy unverified data will fall here and be hidden
            ws.row_dimensions[row_idx].hidden = True
            stats["hidden_unverified"] += 1

    # Hide telemetry metadata columns (Cols 7 through 11, G through K) to preserve clean business view
    for c in range(7, 12):
        col_letter = openpyxl.utils.get_column_letter(c)
        ws.column_dimensions[col_letter].hidden = True

    wb.save(excel_file_path)
    logger.info(f"Failsafe Execution Complete. Stats: {json.dumps(stats, indent=2)}")
    return stats


if __name__ == "__main__":
    from pathlib import Path

    file_path = Path(__file__).resolve().parents[2] / "2027_BCI_Internship_Tracker.xlsx"
    enforce_zero_hallucination_excel_failsafe(str(file_path))
